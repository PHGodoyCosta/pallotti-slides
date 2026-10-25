"""The shipped template (MODELO_PPTX.md) and the inspector rules."""

from collections.abc import Callable
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide

from pallotti_slides.domain.enums import EucharisticPrayer, SongMoment
from pallotti_slides.presentation.errors import TemplateError
from pallotti_slides.presentation.template import (
    MOMENT_TAGS,
    SONG_TEXT_FIELD,
    MassTemplate,
    SlideTag,
    TagKind,
    default_template_path,
)


@pytest.fixture(scope="module")
def template() -> MassTemplate:
    return MassTemplate.load()


def edit_template(tmp_path: Path, change: Callable[[PresentationType], None]) -> MassTemplate:
    """Copy of the shipped template with ``change`` applied, loaded back."""
    presentation = Presentation(str(default_template_path()))
    change(presentation)
    path = tmp_path / "edited.pptx"
    presentation.save(str(path))
    return MassTemplate.load(path)


def first_slide_with(presentation: PresentationType, text: str) -> Slide:
    for slide in presentation.slides:
        if any(s.has_text_frame and text in s.text_frame.text for s in slide.shapes):
            return slide
    raise AssertionError(f"no slide with {text!r}")


def slide_with_tag(presentation: PresentationType, tag: str) -> Slide:
    for slide in presentation.slides:
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text == f"@pallotti {tag}":
            return slide
    raise AssertionError(f"no slide tagged {tag!r}")


def set_notes(slide: Slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def messages(template: MassTemplate, level: str) -> list[str]:
    return [n.message for n in template.validate() if n.level == level]


# --- the shipped template ------------------------------------------------------------------


def test_shipped_template_has_no_problems(template: MassTemplate) -> None:
    assert template.validate() == []


def test_every_slide_is_explicitly_tagged(template: MassTemplate) -> None:
    assert all(slide.tag.explicit for slide in template.slides)
    assert all(slide.tag.kind is not TagKind.INVALID for slide in template.slides)


def test_one_prototype_per_moment_in_mass_order(template: MassTemplate) -> None:
    prototypes = [s.tag.moment for s in template.slides if s.tag.kind is TagKind.SONG]
    assert prototypes == list(SongMoment)


def test_prototypes_have_label_text_and_box(template: MassTemplate) -> None:
    for slide in (s for s in template.slides if s.tag.kind is TagKind.SONG):
        assert {"song.label", SONG_TEXT_FIELD} <= slide.fields
        assert slide.text_box is not None
        assert slide.text_box.width_pt > 600  # wide content box
        assert slide.text_box.height_pt > 300


def test_each_song_moment_has_its_separator(template: MassTemplate) -> None:
    groups = {s.tag.group for s in template.slides if s.tag.kind is TagKind.GROUP}
    wanted = {
        f"canto:{MOMENT_TAGS[m]}"
        for m in (
            SongMoment.ENTRANCE,
            SongMoment.PENITENTIAL_ACT,
            SongMoment.GLORIA,
            SongMoment.OFFERTORY,
            SongMoment.COMMUNION,
            SongMoment.POST_COMMUNION,
            SongMoment.RECESSIONAL,
        )
    }
    assert wanted <= groups


def test_separator_comes_right_before_its_prototype(template: MassTemplate) -> None:
    for index, slide in enumerate(template.slides):
        if slide.tag.kind is TagKind.SONG and index > 0:
            previous = template.slides[index - 1]
            if previous.tag.group and previous.tag.group.startswith("canto:"):
                assert previous.tag.group == slide.tag.group


def test_second_reading_group_holds_the_reading_slides(template: MassTemplate) -> None:
    group = [s for s in template.slides if s.tag.group == "segunda_leitura"]
    assert len(group) == 2
    assert any("liturgy.second_reading" in s.fields for s in group)


def test_prayer_block_is_contiguous_and_has_mystery_of_faith(template: MassTemplate) -> None:
    assert template.available_prayers() == [
        EucharisticPrayer.III,
        EucharisticPrayer.VARIOUS_NEEDS_IV,
    ]
    block = [s.index for s in template.slides if s.tag.prayer is EucharisticPrayer.III]
    assert block == list(range(block[0], block[-1] + 1))
    assert any("liturgy.mystery_of_faith" in template.slides[i].fields for i in block)
    block_dc = [
        s.index for s in template.slides if s.tag.prayer is EucharisticPrayer.VARIOUS_NEEDS_IV
    ]
    assert len(block_dc) == 18  # 9 images + 9 responses
    assert block_dc == list(range(block_dc[0], block_dc[-1] + 1))
    assert any("liturgy.mystery_of_faith" in template.slides[i].fields for i in block_dc)


def test_all_text_fields_are_in_the_template(template: MassTemplate) -> None:
    used = set().union(*(s.fields for s in template.slides))
    assert used == {
        "celebration.title",
        "liturgy.first_reading",
        "liturgy.psalm_number",
        "liturgy.psalm_refrain",
        "liturgy.second_reading",
        "liturgy.gospel",
        "liturgy.intercessions_response",
        "liturgy.mystery_of_faith",
        "song.label",
        "song.text",
    }


def test_tag_keeps_the_formatting_of_the_original_run(tmp_path: Path) -> None:
    presentation = Presentation(str(default_template_path()))
    title = first_slide_with(presentation, "{{celebracao.titulo}}")
    runs = [
        r for s in title.shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs
    ]
    run = next(r for r in runs if r.text == "{{celebracao.titulo}}")
    assert run.font.bold is True
    assert run.font.size is not None
    assert run.font.size.pt == 36


def test_psalm_and_intercessions_keep_literal_prefixes(tmp_path: Path) -> None:
    presentation = Presentation(str(default_template_path()))
    psalm = first_slide_with(presentation, "{{liturgia.salmo_numero}}")
    texts = [p.text for s in psalm.shapes if s.has_text_frame for p in s.text_frame.paragraphs]
    assert "Salmo {{liturgia.salmo_numero}}" in texts
    assert "{{liturgia.salmo_refrao}}" in texts
    intercessions = first_slide_with(presentation, "{{liturgia.preces_refrao}}")
    texts = [
        p.text for s in intercessions.shapes if s.has_text_frame for p in s.text_frame.paragraphs
    ]
    assert "R: {{liturgia.preces_refrao}}" in texts


def test_missing_file_is_a_template_error(tmp_path: Path) -> None:
    with pytest.raises(TemplateError, match="não encontrado"):
        MassTemplate.load(tmp_path / "nope.pptx")


def test_corrupt_file_is_a_template_error(tmp_path: Path) -> None:
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"not a zip")
    with pytest.raises(TemplateError, match="Não foi possível abrir"):
        MassTemplate.load(broken)


# --- slide tag parser ------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("line", "kind"),
    [
        ("@pallotti fixo", TagKind.FIXED),
        ("@pallotti grupo=canto:entrada", TagKind.GROUP),
        ("@pallotti grupo=segunda_leitura", TagKind.GROUP),
        ("@pallotti canto=ato_penitencial", TagKind.SONG),
        ("@pallotti canto=pos_comunhao", TagKind.SONG),
        ("@pallotti oe=III", TagKind.PRAYER),
        ("  @pallotti   fixo  ", TagKind.FIXED),
    ],
)
def test_parse_valid_tags(line: str, kind: TagKind) -> None:
    tag = SlideTag.parse(line)
    assert tag.kind is kind
    assert tag.explicit
    assert tag.error is None


@pytest.mark.parametrize(
    "line",
    [
        "@pallotti",
        "@pallotti foo",
        "@pallotti canto=",
        "@pallotti canto=xyz",
        "@pallotti grupo=xyz",
        "@pallotti grupo=canto:xyz",
        "@pallotti oe=V",
        "@pallotti fixo extra",
        "@pallotti canto=entrada oe=III",
    ],
)
def test_parse_invalid_tags(line: str) -> None:
    tag = SlideTag.parse(line)
    assert tag.kind is TagKind.INVALID
    assert tag.error


@pytest.mark.parametrize("line", [None, "", "  ", "notas qualquer"])
def test_missing_tag_is_implicit_fixed(line: str | None) -> None:
    tag = SlideTag.parse(line)
    assert tag.kind is TagKind.FIXED
    assert not tag.explicit


def test_song_tag_maps_to_moment_and_group() -> None:
    tag = SlideTag.parse("@pallotti canto=pos_comunhao")
    assert tag.moment is SongMoment.POST_COMMUNION
    assert tag.group == "canto:pos_comunhao"
    assert SlideTag.parse("@pallotti oe=IV").prayer is EucharisticPrayer.IV


# --- inspector rules (broken copies of the template) -------------------------------------------


def test_invalid_slide_tag_is_an_error(tmp_path: Path) -> None:
    def change(p: PresentationType) -> None:
        set_notes(p.slides[0], "@pallotti canto=xyz")

    errors = messages(edit_template(tmp_path, change), "error")
    assert any("momento de canto desconhecido" in m for m in errors)


def test_untagged_slide_is_a_warning(tmp_path: Path) -> None:
    def change(p: PresentationType) -> None:
        set_notes(p.slides[0], "")

    template = edit_template(tmp_path, change)
    assert not messages(template, "error")
    assert any("sem etiqueta" in m for m in messages(template, "warning"))


def test_missing_prototype_is_an_error(tmp_path: Path) -> None:
    def change(p: PresentationType) -> None:
        set_notes(slide_with_tag(p, "canto=mantra"), "@pallotti fixo")

    errors = messages(edit_template(tmp_path, change), "error")
    assert any("Falta o protótipo de canto 'mantra'" in m for m in errors)


def test_duplicate_prototype_is_an_error(tmp_path: Path) -> None:
    def change(p: PresentationType) -> None:
        set_notes(slide_with_tag(p, "canto=mantra"), "@pallotti canto=entrada")

    errors = messages(edit_template(tmp_path, change), "error")
    assert any("duplicado" in m for m in errors)
    assert any("Falta o protótipo de canto 'mantra'" in m for m in errors)


def test_split_tag_is_an_error(tmp_path: Path) -> None:
    def change(p: PresentationType) -> None:
        slide = first_slide_with(p, "{{celebracao.titulo}}")
        run = next(
            r
            for s in slide.shapes
            if s.has_text_frame
            for para in s.text_frame.paragraphs
            for r in para.runs
            if r.text == "{{celebracao.titulo}}"
        )
        run.text = "{{celebracao."
        run._r.addnext(_clone_with_text(run, "titulo}}"))

    template = edit_template(tmp_path, change)
    assert any("dividida" in m for m in messages(template, "error"))


def _clone_with_text(run, text: str):  # type: ignore[no-untyped-def]
    from copy import deepcopy

    clone = deepcopy(run._r)
    clone.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = text
    return clone


def test_unknown_text_tag_is_an_error(tmp_path: Path) -> None:
    def change(p: PresentationType) -> None:
        slide = first_slide_with(p, "{{liturgia.evangelho}}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.text == "{{liturgia.evangelho}}":
                            run.text = "{{liturgia.evangelhoo}}"

    errors = messages(edit_template(tmp_path, change), "error")
    assert any("{{liturgia.evangelhoo}}" in m for m in errors)


def test_prototype_without_song_text_is_an_error(tmp_path: Path) -> None:
    def change(p: PresentationType) -> None:
        for slide in p.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if run.text == "{{canto.texto}}":
                                run.text = "texto fixo"
                                return

    errors = messages(edit_template(tmp_path, change), "error")
    assert any("sem {{canto.texto}}" in m for m in errors)


def test_song_tag_outside_prototype_is_an_error(tmp_path: Path) -> None:
    def change(p: PresentationType) -> None:
        slide = first_slide_with(p, "{{liturgia.evangelho}}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.text == "{{liturgia.evangelho}}":
                            run.text = "{{canto.texto}}"

    errors = messages(edit_template(tmp_path, change), "error")
    assert any("só podem ficar em protótipo" in m for m in errors)


def test_prayer_block_without_mystery_of_faith_is_a_warning(tmp_path: Path) -> None:
    def change(p: PresentationType) -> None:
        slide = first_slide_with(p, "{{liturgia.misterio_fe}}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.text == "{{liturgia.misterio_fe}}":
                            run.text = "Salvador do mundo"

    template = edit_template(tmp_path, change)
    assert not messages(template, "error")
    assert any("misterio_fe" in m for m in messages(template, "warning"))


def test_template_without_prayer_block_is_an_error(tmp_path: Path) -> None:
    def change(p: PresentationType) -> None:
        for slide in p.slides:
            if slide.has_notes_slide and "oe=" in slide.notes_slide.notes_text_frame.text:
                set_notes(slide, "@pallotti fixo")

    errors = messages(edit_template(tmp_path, change), "error")
    assert any("nenhum bloco de Oração Eucarística" in m for m in errors)
