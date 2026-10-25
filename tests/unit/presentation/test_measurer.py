"""TextMeasurer: calibration against the example deck (SPEC F9) and basic behaviour."""

import re
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.shapes.autoshape import Shape

from pallotti_slides.domain.models import Line
from pallotti_slides.presentation.measurer import (
    EMU_PER_POINT,
    SUBSTITUTE_SLACK,
    BodyFormat,
    Box,
    TextMeasurer,
)
from pallotti_slides.presentation.template import read_body_format

ROOT = Path(__file__).resolve().parents[3]
EXAMPLE_PPTX = ROOT / "examples" / "Missa - 26° Dom do T.C - 26.09.2026 - OE III.pptx"
EXAMPLE_PDF = EXAMPLE_PPTX.with_suffix(".pdf")

# Song slides of the example (1-based): label in the top-left corner + one body.
SONG_SLIDES = (3, 5, 6, 7, 8, 10, 11, 12, 13, 15, 16, 17, 18, 19, 26, 36, 37, 38, 39)
SONG_LABELS = {"Mantra", "Canto de Entrada", "Ato Penitencial", "Glória", "Aclamação", "Ofertório"}


@dataclass(frozen=True)
class ExampleSlide:
    number: int
    box: Box
    body: BodyFormat
    lines: list[Line]
    size_pt: float


def _example_song_slides() -> list[ExampleSlide]:
    presentation = Presentation(str(EXAMPLE_PPTX))
    slides = list(presentation.slides)
    found: list[ExampleSlide] = []
    for number in SONG_SLIDES:
        shapes = [
            s
            for s in slides[number - 1].shapes
            if isinstance(s, Shape) and s.has_text_frame and s.text_frame.text.strip()
        ]
        assert any(s.text_frame.text.strip() in SONG_LABELS for s in shapes), number
        (body_shape,) = [s for s in shapes if s.text_frame.text.strip() not in SONG_LABELS]
        paragraphs = [p for p in body_shape.text_frame.paragraphs if p.text.strip()]
        sizes = {run.font.size.pt for p in paragraphs for run in p.runs if run.font.size}
        assert len(sizes) == 1, number
        lines = [Line(text=p.text, bold=bool(p.runs[0].font.bold)) for p in paragraphs]
        box = Box(
            int(body_shape.left or 0),
            int(body_shape.top or 0),
            int(body_shape.width or 0),
            int(body_shape.height or 0),
        )
        body = read_body_format(body_shape.text_frame._txBody, paragraphs[0]._p)
        found.append(ExampleSlide(number, box, body, lines, sizes.pop()))
    return found


EXAMPLE_SLIDES = _example_song_slides()
BODY_BOTTOM_MARGIN_EMU = 182_880  # 0.2"
SLIDE_HEIGHT = int(Presentation(str(EXAMPLE_PPTX)).slide_height or 0)


def _box_to_the_bottom(box: Box) -> Box:
    """The example box stretched down to 0.2" above the slide bottom, as in the template."""
    height = SLIDE_HEIGHT - BODY_BOTTOM_MARGIN_EMU - box.top
    return Box(box.left, box.top, box.width, max(box.height, height))


@pytest.mark.parametrize("slide", EXAMPLE_SLIDES, ids=lambda s: f"slide{s.number}")
def test_calibration_every_song_slide_of_the_example_fits(
    measurer: TextMeasurer, slide: ExampleSlide
) -> None:
    area = _box_to_the_bottom(slide.box)
    assert measurer.fits(slide.lines, slide.size_pt, area, slide.body)


def _pdf_body_lines(page: int) -> list[tuple[float, float]]:
    """(yMin, yMax) of the text lines of a PDF page, without the song label at the top."""
    result = subprocess.run(
        ["pdftotext", "-f", str(page), "-l", str(page), "-bbox-layout", str(EXAMPLE_PDF), "-"],
        capture_output=True,
        text=True,
        check=True,
    )
    pattern = r'<line xMin="[\d.]+" yMin="([\d.]+)" xMax="[\d.]+" yMax="([\d.]+)">'
    lines = [(float(a), float(b)) for a, b in re.findall(pattern, result.stdout)]
    return [line for line in lines if line[0] > 57]  # the label ends at y = 56 pt


@pytest.mark.skipif(shutil.which("pdftotext") is None, reason="precisa do pdftotext (poppler)")
@pytest.mark.parametrize("slide", EXAMPLE_SLIDES, ids=lambda s: f"slide{s.number}")
def test_calibration_matches_the_libreoffice_rendering(
    measurer: TextMeasurer, slide: ExampleSlide
) -> None:
    rendered = _pdf_body_lines(slide.number)
    wrapped = sum(
        measurer.line_count(line, slide.size_pt, slide.box, slide.body) for line in slide.lines
    )
    assert wrapped == len(rendered)
    top = slide.box.top / EMU_PER_POINT + slide.body.inset_top_pt
    bottom = top + measurer.height(slide.lines, slide.size_pt, slide.box, slide.body)
    assert bottom == pytest.approx(max(y for _, y in rendered), abs=1.0)


def test_example_slides_12_and_38_run_below_their_box(measurer: TextMeasurer) -> None:
    """Documented: in the example these two slides pass the body box (not the slide)."""
    by_number = {s.number: s for s in EXAMPLE_SLIDES}
    for number in (12, 38):
        slide = by_number[number]
        assert not measurer.fits(slide.lines, slide.size_pt, slide.box, slide.body)
    for slide in EXAMPLE_SLIDES:
        if slide.number not in (12, 38):
            assert measurer.fits(slide.lines, slide.size_pt, slide.box, slide.body)


def test_wrap_and_height(measurer: TextMeasurer) -> None:
    box = Box(0, 0, 400 * EMU_PER_POINT, 400 * EMU_PER_POINT)
    body = BodyFormat(margin_left_pt=0, indent_pt=0)
    short = [Line(text="Curta")]
    long = [Line(text="palavra " * 30)]
    assert measurer.line_count(short[0], 40, box, body) == 1
    assert measurer.line_count(long[0], 40, box, body) > 3
    assert measurer.height(long, 40, box, body) > measurer.height(short, 40, box, body)
    assert measurer.height([], 40, box, body) == 0
    assert measurer.fits(short, 40, box, body)
    assert not measurer.fits(long * 3, 40, box, body)


def test_bold_is_wider(measurer: TextMeasurer) -> None:
    assert measurer.text_width("Glória", True, 54) > measurer.text_width("Glória", False, 54)


def test_word_longer_than_the_line_counts_several_lines(measurer: TextMeasurer) -> None:
    box = Box(0, 0, 100 * EMU_PER_POINT, 400 * EMU_PER_POINT)
    body = BodyFormat(margin_left_pt=0, indent_pt=0)
    assert measurer.line_count(Line(text="X" * 40), 40, box, body) >= 3


def test_paragraph_spacing_comes_from_the_body_format(measurer: TextMeasurer) -> None:
    box = Box(0, 0, 900 * EMU_PER_POINT, 500 * EMU_PER_POINT)
    lines = [Line(text="Um"), Line(text="Dois")]
    tight = BodyFormat(space_before_pt=0, space_after_pt=0)
    loose = BodyFormat(space_before_pt=10, space_after_pt=5)
    assert measurer.height(lines, 40, box, loose) == pytest.approx(
        measurer.height(lines, 40, box, tight) + 15
    )


def test_without_aptos_uses_a_substitute_with_slack_and_a_notice(tmp_path: Path) -> None:
    measurer = TextMeasurer(tmp_path / "missing.ttf", None)
    assert measurer.uses_substitute
    (notice,) = measurer.notices
    assert notice.level == "warning"
    assert "Aptos" in notice.message
    box = Box(0, 0, 900 * EMU_PER_POINT, 500 * EMU_PER_POINT)
    exact = BodyFormat(line_spacing_pt=50, space_before_pt=0, space_after_pt=0)
    height = measurer.height([Line(text="Um")], 40, box, exact)
    assert height == pytest.approx((50 + 0.06 * 40) * SUBSTITUTE_SLACK)
    assert measurer.text_width("abc", False, 40) > 0


def test_with_aptos_there_is_no_notice(measurer: TextMeasurer) -> None:
    assert not measurer.uses_substitute
    assert measurer.notices == []
