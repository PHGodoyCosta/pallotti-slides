"""Duplicate / keep / reorder slides at the XML level."""

import importlib.util
import io
import zipfile
from pathlib import Path
from types import ModuleType

import pytest
from pptx import Presentation
from pptx.slide import Slide

from pallotti_slides.presentation.slides_xml import (
    SlideImporter,
    SlideImportError,
    duplicate_slide,
    keep_slides,
)
from pallotti_slides.presentation.template import default_prayers_dir, default_template_path

ROOT = Path(__file__).resolve().parents[3]
PICTURE = 13


def load_validator() -> ModuleType:
    spec = importlib.util.spec_from_file_location(
        "validate_ooxml", ROOT / "scripts" / "validate_ooxml.py"
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def pictures(slide: Slide) -> list[str]:
    return [s.image.sha1 for s in slide.shapes if s.shape_type == PICTURE]  # type: ignore[attr-defined]


def notes(slide: Slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""  # type: ignore[union-attr]


def reload(presentation: Presentation) -> Presentation:  # type: ignore[valid-type]
    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)
    return Presentation(buffer)


def test_duplicate_slide_copies_content_and_shares_the_image() -> None:
    presentation = Presentation(str(default_template_path()))
    source = presentation.slides[3]  # separator image
    total = len(presentation.slides)
    copy = duplicate_slide(presentation, source)

    assert len(presentation.slides) == total + 1
    assert presentation.slides[-1].slide_id == copy.slide_id
    assert copy.slide_layout.name == source.slide_layout.name
    source_pics = [s for s in source.shapes if s.shape_type == 13]
    copy_pics = [s for s in copy.shapes if s.shape_type == 13]
    assert len(copy_pics) == len(source_pics) == 1
    assert copy_pics[0].image.sha1 == source_pics[0].image.sha1  # type: ignore[attr-defined]


def test_duplicate_slide_survives_save_and_reload() -> None:
    presentation = Presentation(str(default_template_path()))
    prototype = next(
        s
        for s in presentation.slides
        if any(sh.has_text_frame and "{{canto.texto}}" in sh.text_frame.text for sh in s.shapes)
    )
    duplicate_slide(presentation, prototype)
    reloaded = reload(presentation)
    texts = [
        sh.text_frame.text
        for sh in reloaded.slides[-1].shapes
        if sh.has_text_frame and sh.text_frame.text
    ]
    assert "{{canto.texto}}" in texts
    assert "{{canto.rotulo}}" in texts


def test_keep_slides_removes_and_reorders() -> None:
    presentation = Presentation(str(default_template_path()))
    slides = list(presentation.slides)
    kept = [slides[5], slides[1], slides[0]]
    keep_slides(presentation, kept)

    reloaded = reload(presentation)
    assert len(reloaded.slides) == 3
    assert [s.slide_id for s in presentation.slides] == [s.slide_id for s in kept]


def test_keep_slides_drops_orphan_parts() -> None:
    presentation = Presentation(str(default_template_path()))
    keep_slides(presentation, [presentation.slides[0]])
    buffer = io.BytesIO()
    presentation.save(buffer)
    with zipfile.ZipFile(buffer) as archive:
        names = archive.namelist()
    assert [n for n in names if n.startswith("ppt/slides/slide")] == ["ppt/slides/slide1.xml"]
    assert len([n for n in names if n.startswith("ppt/notesSlides/")]) <= 2  # notes + rels


# --- slides of another file (M3b) ---------------------------------------------------------------


def test_import_slides_of_another_file(tmp_path: Path) -> None:
    deck = Presentation(str(default_template_path()))
    keep_slides(deck, [deck.slides[0]])
    block = Presentation(str(default_prayers_dir() / "III.pptx"))
    importer = SlideImporter(deck)
    copies = [importer.import_slide(slide) for slide in block.slides]

    assert len(deck.slides) == 1 + len(block.slides)
    for source, copy in zip(block.slides, copies, strict=True):
        assert copy.slide_layout.name == source.slide_layout.name
        assert copy.slide_layout.part.partname == source.slide_layout.part.partname
        assert pictures(copy) == pictures(source)
        assert notes(copy) == notes(source)
    assert any(pictures(copy) for copy in copies)

    path = tmp_path / "imported.pptx"
    deck.save(str(path))
    assert load_validator().validate(path, allow_tags=True) == []
    reloaded = Presentation(str(path))
    assert [pictures(s) for s in list(reloaded.slides)[1:]] == [pictures(s) for s in block.slides]


def test_imported_images_are_deduplicated(tmp_path: Path) -> None:
    deck = Presentation(str(default_template_path()))
    keep_slides(deck, [deck.slides[0]])
    block = Presentation(str(default_prayers_dir() / "DC_IV.pptx"))
    importer = SlideImporter(deck)
    for slide in block.slides:  # 9 copies of the same host image
        importer.import_slide(slide)
    SlideImporter(deck).import_slide(block.slides[0])  # a new importer reuses the part too
    path = tmp_path / "dedup.pptx"
    deck.save(str(path))
    with zipfile.ZipFile(path) as archive:
        media = [n for n in archive.namelist() if n.startswith("ppt/media/")]
    images = {sha for s in Presentation(str(path)).slides for sha in pictures(s)}
    assert len(images) >= 1
    # every image is stored once, plus the images of the master/layouts
    assert len(media) == len(set(media))
    with zipfile.ZipFile(path) as archive:
        blobs = [archive.read(n) for n in media]
    assert len(blobs) == len(set(blobs))


def test_import_needs_an_equivalent_layout() -> None:
    deck = Presentation(str(default_template_path()))
    block = Presentation(str(default_prayers_dir() / "III.pptx"))
    source = block.slides[0]
    source.slide_layout.name = "Layout que não existe"  # type: ignore[misc]
    with pytest.raises(SlideImportError, match="layout"):
        SlideImporter(deck).import_slide(source)
