"""SongDeckCache: one .pptx per song, regenerated when lyrics, typography or prototype change."""

from pathlib import Path

import pytest
from pptx import Presentation

from pallotti_slides.config import TypographyConfig
from pallotti_slides.domain.enums import SongMoment
from pallotti_slides.domain.models import Line, Mass, Song
from pallotti_slides.presentation.measurer import TextMeasurer
from pallotti_slides.presentation.paginator import SongPaginator
from pallotti_slides.presentation.song_cache import SongDeckCache
from pallotti_slides.presentation.template import MassTemplate


def make_cache(
    folder: Path,
    template: MassTemplate,
    measurer: TextMeasurer,
    typography: TypographyConfig | None = None,
) -> SongDeckCache:
    return SongDeckCache(folder, template, SongPaginator(measurer, typography))


def texts(path: Path) -> list[list[str]]:
    return [
        [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text]  # type: ignore[attr-defined]
        for slide in Presentation(str(path)).slides
    ]


@pytest.fixture
def song(mass_26tc: Mass) -> Song:
    return mass_26tc.songs[SongMoment.ENTRANCE]


def test_deck_has_one_slide_per_page_and_keeps_the_label_tag(
    tmp_path: Path, template: MassTemplate, measurer: TextMeasurer, song: Song
) -> None:
    cache = make_cache(tmp_path, template, measurer)
    deck = cache.deck(SongMoment.ENTRANCE, song)
    assert deck is not None
    assert deck.path.parent == tmp_path and deck.path.is_file()
    pages = cache._paginator.paginate(
        song, template.song_area(SongMoment.ENTRANCE), template.body_format(SongMoment.ENTRANCE)
    )
    slides = texts(deck.path)
    assert deck.pages == len(slides) == len(pages) == 4
    for slide, page in zip(slides, pages, strict=True):
        assert "{{canto.rotulo}}" in slide
        assert "{{canto.texto}}" not in " ".join(slide)
        assert "\n".join(line.text for line in page.lines) in slide
    notes = [s.notes_slide.notes_text_frame.text for s in Presentation(str(deck.path)).slides]
    assert all("@pallotti" not in n for n in notes)


def test_second_use_does_not_regenerate(
    tmp_path: Path, template: MassTemplate, measurer: TextMeasurer, song: Song
) -> None:
    first = make_cache(tmp_path, template, measurer)
    deck = first.deck(SongMoment.ENTRANCE, song)
    assert deck is not None and first.generated == [deck.path]
    written = deck.path.stat().st_mtime_ns

    second = make_cache(tmp_path, template, measurer)  # e.g. the next run of the program
    again = second.deck(SongMoment.ENTRANCE, song)
    assert again == deck
    assert second.generated == []
    assert deck.path.stat().st_mtime_ns == written


def test_changing_the_lyrics_regenerates(
    tmp_path: Path, template: MassTemplate, measurer: TextMeasurer, song: Song
) -> None:
    cache = make_cache(tmp_path, template, measurer)
    old = cache.deck(SongMoment.ENTRANCE, song)
    stanzas = list(song.stanzas)
    first = stanzas[0]
    stanzas[0] = first.model_copy(update={"lines": [*first.lines[:-1], Line(text="Nova linha")]})
    new = cache.deck(SongMoment.ENTRANCE, song.model_copy(update={"stanzas": stanzas}))
    assert old is not None and new is not None
    assert new.path != old.path
    assert cache.generated == [old.path, new.path]
    assert "Nova linha" in " ".join(texts(new.path)[0])


def test_changing_the_typography_regenerates(
    tmp_path: Path, template: MassTemplate, measurer: TextMeasurer, song: Song
) -> None:
    default = make_cache(tmp_path, template, measurer).deck(SongMoment.ENTRANCE, song)
    smaller = make_cache(tmp_path, template, measurer, TypographyConfig(min_font=30, max_font=40))
    deck = smaller.deck(SongMoment.ENTRANCE, song)
    assert default is not None and deck is not None
    assert deck.path != default.path
    assert smaller.generated == [deck.path]


def test_same_prototype_shares_the_file_other_prototype_does_not(
    tmp_path: Path, template: MassTemplate, measurer: TextMeasurer, song: Song
) -> None:
    cache = make_cache(tmp_path, template, measurer)
    communion = cache.deck(SongMoment.COMMUNION, song)
    recessional = cache.deck(SongMoment.RECESSIONAL, song)  # same prototype (copy) as communion
    assert communion is not None and recessional is not None
    assert recessional.path == communion.path
    assert cache.generated == [communion.path]
    entrance = cache.deck(SongMoment.ENTRANCE, song)  # narrower body box: its own deck
    assert entrance is not None and entrance.path != communion.path


def test_broken_cache_entry_is_regenerated(
    tmp_path: Path, template: MassTemplate, measurer: TextMeasurer, song: Song
) -> None:
    cache = make_cache(tmp_path, template, measurer)
    deck = cache.deck(SongMoment.ENTRANCE, song)
    assert deck is not None
    deck.path.with_suffix(".json").write_text("{", encoding="utf-8")
    again = make_cache(tmp_path, template, measurer)
    assert again.deck(SongMoment.ENTRANCE, song) == deck
    assert again.generated == [deck.path]


def test_song_without_lines_has_no_deck(
    tmp_path: Path, template: MassTemplate, measurer: TextMeasurer, song: Song
) -> None:
    empty = song.model_copy(
        update={"stanzas": [s.model_copy(update={"lines": []}) for s in song.stanzas]}
    )
    assert make_cache(tmp_path, template, measurer).deck(SongMoment.ENTRANCE, empty) is None
    assert list(tmp_path.iterdir()) == []
