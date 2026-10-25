"""Reference test (SPEC F8): 26tc_2026.json -> same slides and texts as the example deck."""

import importlib.util
from pathlib import Path
from types import ModuleType

import pytest
from pptx import Presentation
from pptx.slide import Slide

from pallotti_slides.application.generate_mass import (
    GenerationResult,
    load_mass,
    make_mass_generator,
)
from pallotti_slides.config import FontLocator
from pallotti_slides.domain.enums import EucharisticPrayer
from pallotti_slides.presentation.preview import pdf_page_count

ROOT = Path(__file__).resolve().parents[2]
EXAMPLE = ROOT / "examples" / "Missa - 26° Dom do T.C - 26.09.2026 - OE III.pptx"
FIXTURE = ROOT / "tests" / "fixtures" / "masses" / "26tc_2026.json"

needs_aptos = pytest.mark.skipif(
    not FontLocator().locate().complete,
    reason="Aptos não instalada: sem ela a paginação pode diferir do exemplo",
)


def load_validator() -> ModuleType:
    spec = importlib.util.spec_from_file_location(
        "validate_ooxml", ROOT / "scripts" / "validate_ooxml.py"
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def normalize(text: str) -> str:
    return " ".join(text.split())


def slide_texts(slide: Slide) -> list[str]:
    """Non-empty paragraphs of every text shape, normalized and sorted (shape order aside)."""
    found = [
        normalize(paragraph.text)
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs  # type: ignore[attr-defined]
        if normalize(paragraph.text)
    ]
    return sorted(found)


def slide_pictures(slide: Slide) -> list[str]:
    return sorted(s.image.sha1 for s in slide.shapes if s.shape_type == 13)  # type: ignore[attr-defined]


@pytest.fixture(scope="module")
def generated(tmp_path_factory: pytest.TempPathFactory) -> GenerationResult:
    output = tmp_path_factory.mktemp("example")
    return make_mass_generator(with_pdf=False).generate(load_mass(FIXTURE), output)


@needs_aptos
def test_same_slide_sequence_and_texts_as_the_example(generated: GenerationResult) -> None:
    ours = list(Presentation(str(generated.pptx)).slides)
    theirs = list(Presentation(str(EXAMPLE)).slides)
    assert len(ours) == len(theirs) == 69
    for number, (mine, reference) in enumerate(zip(ours, theirs, strict=True), start=1):
        assert slide_texts(mine) == slide_texts(reference), f"slide {number}"
        assert slide_pictures(mine) == slide_pictures(reference), f"slide {number}"


def test_file_name_and_structure(generated: GenerationResult) -> None:
    assert generated.pptx.name == "Missa - 26° Dom do T.C - 26.09.2026 - OE III.pptx"
    assert generated.pdf is None
    assert load_validator().validate(generated.pptx) == []


def test_no_tag_left_in_slides_or_notes(generated: GenerationResult) -> None:
    deck = Presentation(str(generated.pptx))
    for slide in deck.slides:
        text = " ".join(slide_texts(slide))
        assert "{{" not in text and "@pallotti" not in text
        if slide.has_notes_slide:
            assert "@pallotti" not in slide.notes_slide.notes_text_frame.text  # type: ignore[union-attr]


@pytest.mark.soffice
def test_real_pdf_has_one_page_per_slide(tmp_path: Path) -> None:
    result = make_mass_generator().generate(load_mass(FIXTURE), tmp_path)
    assert result.pdf is not None and result.pdf.is_file(), result.notices
    assert result.pdf.parent == tmp_path
    slides = len(Presentation(str(result.pptx)).slides)
    assert pdf_page_count(result.pdf) == slides == 69


def test_various_needs_iv_responses_in_order(tmp_path: Path) -> None:
    mass = load_mass(FIXTURE)
    liturgy = mass.liturgy.model_copy(
        update={"eucharistic_prayer": EucharisticPrayer.VARIOUS_NEEDS_IV}
    )
    mass = mass.model_copy(update={"liturgy": liturgy})
    result = make_mass_generator(with_pdf=False).generate(mass, tmp_path)
    assert result.pptx.name.endswith("OE DC IV.pptx")
    texts = [slide_texts(s) for s in Presentation(str(result.pptx)).slides]
    start = next(
        i for i, t in enumerate(texts) if t == ["R: Bendito o vosso Filho, presente entre nós!"]
    )
    responses = [t[0] for t in texts[start : start + 18] if t]
    assert responses == [
        "R: Bendito o vosso Filho, presente entre nós!",
        "R: Enviai o vosso Espírito Santo!",
        mass.liturgy.mystery_of_faith,
        "R: Aceitai, ó Senhor, a nossa oferta!",
        "R: O Espírito nos una num só corpo!",
        "R: Confirmai na unidade a vossa Igreja!",
        "R: Ajudai-nos a criar um mundo novo!",
        "R: Concedei-lhes, ó Senhor, a luz eterna!",
        "R: Amém!",
    ]
