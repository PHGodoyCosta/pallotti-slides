"""End to end with fakes: leaflet (FakeLLMClient) + pasted lyrics -> pptx/pdf + history."""

import datetime as dt
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from pallotti_slides.application.container import Container, create_container
from pallotti_slides.config import FontLocator, Paths, Settings, locate_soffice
from pallotti_slides.domain.enums import EucharisticPrayer, SongMoment
from pallotti_slides.domain.models import SongRequest
from pallotti_slides.llm.client import FakeLLMClient

ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "tests" / "fixtures"
LEAFLET = FIXTURES / "leaflets" / "30tc_2026.pdf"

pytestmark = pytest.mark.skipif(not FontLocator().locate().complete, reason="Aptos não instalada")

LYRICS = """Ó Senhor, vem nos salvar
Nós queremos te louvar

Tua palavra é luz
Que nos conduz ao bem

Ó Senhor, vem nos salvar
Nós queremos te louvar
"""


@pytest.fixture
def llm_dir(tmp_path: Path) -> Path:
    directory = tmp_path / "llm"
    shutil.copytree(FIXTURES / "llm", directory)
    (directory / "ChorusIds.json").write_text('{"chorus_ids": [1]}', encoding="utf-8")
    return directory


@pytest.fixture
def container(tmp_path: Path, llm_dir: Path) -> Container:
    paths = Paths(data=tmp_path / "data", config=tmp_path / "config.json", logs=tmp_path / "logs")
    with_pdf = locate_soffice() is not None
    return create_container(
        Settings(),
        paths=paths,
        db_url=f"sqlite:///{(tmp_path / 'flow.db').as_posix()}",
        with_pdf=with_pdf,
        llm=FakeLLMClient(llm_dir),
    )


def slide_texts(pptx: Path) -> list[list[str]]:
    return [
        [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]  # type: ignore[attr-defined]
        for slide in Presentation(str(pptx)).slides
    ]


def test_full_flow(container: Container, tmp_path: Path) -> None:
    requests = [
        SongRequest(moment=SongMoment.ENTRANCE, text=LYRICS, kind="lyrics"),
        SongRequest(moment=SongMoment.COMMUNION, text="me chamaste", kind="title"),
    ]
    draft = container.preparer.prepare(
        dt.date(2026, 10, 25),
        LEAFLET,
        requests,
        None,
    )
    assert draft.leaflet is not None
    assert draft.leaflet.gospel == "Mt 22,34-40"
    assert SongMoment.ENTRANCE in draft.songs
    assert SongMoment.COMMUNION not in draft.songs  # title search arrives in M7
    assert any(n.field == "songs.communion" for n in draft.notices)

    mass = draft.to_mass(EucharisticPrayer.III)
    out = tmp_path / "saida"
    result = container.generator.generate(mass, out)

    assert result.pptx.is_file()
    if locate_soffice() is not None:
        assert result.pdf is not None and result.pdf.is_file()
    else:
        assert result.pdf is None

    rows = container.mass_repo.list()
    assert len(rows) == 1
    assert rows[0].pptx_path == result.pptx
    assert rows[0].pdf_path == result.pdf
    assert container.mass_repo.get(rows[0].id) == mass


def test_reopen_and_regenerate_gives_same_content(container: Container, tmp_path: Path) -> None:
    request = SongRequest(moment=SongMoment.ENTRANCE, text=LYRICS, kind="lyrics")
    mass = container.preparer.prepare(dt.date(2026, 10, 25), LEAFLET, [request], None).to_mass(
        EucharisticPrayer.III
    )
    first = container.generator.generate(mass, tmp_path / "a")

    reopened = container.mass_repo.get(container.mass_repo.list()[0].id)
    second = container.generator.generate(reopened, tmp_path / "b")

    assert slide_texts(first.pptx) == slide_texts(second.pptx)
    assert len(container.mass_repo.list()) == 2


def test_leaflet_failure_does_not_block_songs(container: Container, tmp_path: Path) -> None:
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"not a pdf")
    request = SongRequest(moment=SongMoment.ENTRANCE, text=LYRICS, kind="lyrics")
    draft = container.preparer.prepare(dt.date(2026, 10, 25), bad, [request], None)
    assert draft.leaflet is None
    assert any(n.level == "error" and n.field == "leaflet" for n in draft.notices)
    assert SongMoment.ENTRANCE in draft.songs
