"""MassGenerator and ``pallotti generate`` with fakes (no LibreOffice needed)."""

from pathlib import Path

import pytest
from pptx import Presentation
from typer.testing import CliRunner

from pallotti_slides.application.generate_mass import (
    MassGenerator,
    load_mass,
    make_mass_generator,
)
from pallotti_slides.cli import app
from pallotti_slides.config import FontLocator, TypographyConfig
from pallotti_slides.db.repositories import MassFiles
from pallotti_slides.domain.enums import SongMoment
from pallotti_slides.domain.errors import PallottiError
from pallotti_slides.domain.models import Mass
from pallotti_slides.presentation.builder import PptxBuilder
from pallotti_slides.presentation.errors import PdfConversionError
from pallotti_slides.presentation.measurer import TextMeasurer
from pallotti_slides.presentation.paginator import SongPaginator
from pallotti_slides.presentation.planner import BuildPlanner
from pallotti_slides.presentation.song_cache import SongDeckCache
from pallotti_slides.presentation.template import MassTemplate

ROOT = Path(__file__).resolve().parents[2]
MASSES = ROOT / "tests" / "fixtures" / "masses"
FULL = MASSES / "full_9_songs.json"
LABELS_IN_ORDER = [m.label for m in SongMoment if m is not SongMoment.GLORIA]

needs_aptos = pytest.mark.skipif(not FontLocator().locate().complete, reason="Aptos não instalada")


class FailingConverter:
    def convert(self, pptx: Path, folder: Path) -> Path:
        raise PdfConversionError("LibreOffice não encontrado.")


class FakeConverter:
    def convert(self, pptx: Path, folder: Path) -> Path:
        pdf = folder / f"{pptx.stem}.pdf"
        pdf.write_bytes(b"%PDF-fake")
        return pdf


class FakeHistory:
    def __init__(self) -> None:
        self.saved: list[tuple[Mass, MassFiles | None]] = []

    def save(self, mass: Mass, files: MassFiles | None = None) -> int:
        self.saved.append((mass, files))
        return len(self.saved)


def generator(
    converter: object,
    history: FakeHistory | None = None,
    songs_dir: Path | None = None,
    typography: TypographyConfig | None = None,
) -> MassGenerator:
    fonts = FontLocator().locate()
    template = MassTemplate.load()
    paginator = SongPaginator(TextMeasurer(fonts.regular, fonts.bold), typography)
    return MassGenerator(
        template=template,
        paginator=paginator,
        planner=BuildPlanner(),
        builder=PptxBuilder(),
        converter=converter,  # type: ignore[arg-type]
        mass_repo=history,
        song_cache=SongDeckCache(songs_dir, template, paginator) if songs_dir else None,
    )


def cache_of(gen: MassGenerator) -> SongDeckCache:
    cache = gen._song_cache
    assert cache is not None
    return cache


def deck_texts(pptx: Path) -> list[list[str]]:
    return [
        sorted(s.text_frame.text for s in slide.shapes if s.has_text_frame)  # type: ignore[attr-defined]
        for slide in Presentation(str(pptx)).slides
    ]


def labels(pptx: Path) -> list[str]:
    """Song label of each song slide, in deck order."""
    found = []
    for slide in Presentation(str(pptx)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text in {m.label for m in SongMoment}:  # type: ignore[attr-defined]
                found.append(shape.text_frame.text)  # type: ignore[attr-defined]
    return found


@needs_aptos
def test_full_mass_songs_in_mass_order(tmp_path: Path) -> None:
    result = generator(FakeConverter()).generate(load_mass(FULL), tmp_path)
    found = labels(result.pptx)
    first_seen = list(dict.fromkeys(found))
    assert first_seen == LABELS_IN_ORDER  # SPEC §3, without the Gloria (has_gloria = false)
    assert found.count("Comunhão") == 3  # 12 lines -> 6 + 6, plus the chorus
    texts = [
        shape.text_frame.text  # type: ignore[attr-defined]
        for slide in Presentation(str(result.pptx)).slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert "Fl 2,1-11" not in " ".join(texts)  # no second reading
    assert result.pdf is not None and result.pdf.is_file()
    assert result.notices == []


@needs_aptos
def test_song_decks_are_cached_between_generations(tmp_path: Path) -> None:
    songs = tmp_path / "songs"
    mass = load_mass(FULL)
    used = len(mass.songs) - 1  # has_gloria = false: the Gloria is not used
    first = generator(None, songs_dir=songs)
    result = first.generate(mass, tmp_path / "a")
    assert len(cache_of(first).generated) == used  # one deck per song
    assert sorted(songs.glob("*.pptx")) == sorted(cache_of(first).generated)

    second = generator(None, songs_dir=songs)  # e.g. the program opened again
    again = second.generate(mass, tmp_path / "b")
    assert cache_of(second).generated == []
    assert deck_texts(again.pptx) == deck_texts(result.pptx)

    smaller = generator(None, songs_dir=songs, typography=TypographyConfig(max_font=44))
    smaller.generate(mass, tmp_path / "c")
    assert len(cache_of(smaller).generated) == used  # typography -> new decks


@needs_aptos
def test_same_song_in_another_moment_reuses_the_deck_with_its_label(tmp_path: Path) -> None:
    songs = tmp_path / "songs"
    mass = load_mass(FULL)
    generator(None, songs_dir=songs).generate(mass, tmp_path / "a")
    communion = mass.songs[SongMoment.COMMUNION]
    moved = mass.model_copy(update={"songs": {**mass.songs, SongMoment.RECESSIONAL: communion}})
    gen = generator(None, songs_dir=songs)
    result = gen.generate(moved, tmp_path / "b")
    assert cache_of(gen).generated == []  # same prototype: the communion deck is reused
    found = labels(result.pptx)
    assert found.count("Comunhão") == found.count("Canto de Saída") == 3
    texts = deck_texts(result.pptx)
    bodies = [[t for t in slide if t not in {m.label for m in SongMoment}] for slide in texts]
    label_of = [
        next((t for t in slide if t in {m.label for m in SongMoment}), None) for slide in texts
    ]
    communion_bodies = [b for b, label in zip(bodies, label_of, strict=True) if label == "Comunhão"]
    recessional_bodies = [
        b for b, label in zip(bodies, label_of, strict=True) if label == "Canto de Saída"
    ]
    assert communion_bodies == recessional_bodies


def test_without_libreoffice_the_pptx_is_still_generated(tmp_path: Path) -> None:
    result = generator(FailingConverter()).generate(load_mass(FULL), tmp_path)
    assert result.pptx.is_file()
    assert result.pdf is None
    warning = next(n for n in result.notices if n.level == "warning" and "PDF" in n.message)
    assert "LibreOffice" in warning.message and "Exportar" in warning.message


def test_history_receives_the_files(tmp_path: Path) -> None:
    history = FakeHistory()
    result = generator(FakeConverter(), history).generate(load_mass(FULL), tmp_path)
    ((mass, files),) = history.saved
    assert mass.celebration.title == "26° Domingo do Tempo Comum"
    assert files == MassFiles(pptx=result.pptx, pdf=result.pdf)


def test_invalid_mass_is_a_fatal_error(tmp_path: Path) -> None:
    mass = load_mass(FULL)
    liturgy = mass.liturgy.model_copy(update={"mystery_of_faith": " "})
    with pytest.raises(PallottiError, match="mistério da fé"):
        generator(FakeConverter()).generate(mass.model_copy(update={"liturgy": liturgy}), tmp_path)
    assert not list(tmp_path.iterdir())


def test_load_mass_errors(tmp_path: Path) -> None:
    bad = tmp_path / "bad.json"
    bad.write_text("{}", encoding="utf-8")
    with pytest.raises(PallottiError, match="inválido"):
        load_mass(bad)
    with pytest.raises(PallottiError, match="Não foi possível ler"):
        load_mass(tmp_path / "missing.json")


def test_make_mass_generator_without_pdf(tmp_path: Path) -> None:
    result = make_mass_generator(with_pdf=False).generate(load_mass(FULL), tmp_path)
    assert result.pdf is None
    assert any(n.level == "info" for n in result.notices)


runner = CliRunner()


def test_cli_generate_without_pdf(tmp_path: Path) -> None:
    res = runner.invoke(
        app, ["generate", "--mass", str(FULL), "--output", str(tmp_path), "--no-pdf"]
    )
    assert res.exit_code == 0, res.output
    assert "PowerPoint:" in res.output
    assert (tmp_path / "Missa - 26° Dom do T.C - 26.09.2026 - OE III.pptx").is_file()


def test_cli_generate_bad_file(tmp_path: Path) -> None:
    bad = tmp_path / "bad.json"
    bad.write_text("not json", encoding="utf-8")
    res = runner.invoke(app, ["generate", "--mass", str(bad), "--output", str(tmp_path / "o")])
    assert res.exit_code == 1
    assert "ERRO" in res.output
