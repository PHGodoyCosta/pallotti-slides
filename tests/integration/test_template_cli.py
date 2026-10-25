"""``pallotti template inspect`` / ``preview`` end to end."""

from pathlib import Path

import pytest
from pptx import Presentation
from typer.testing import CliRunner

from pallotti_slides.application.template_tools import TemplateTools
from pallotti_slides.cli import app
from pallotti_slides.presentation.template import MassTemplate, default_template_path

runner = CliRunner()


def test_inspect_shipped_template_exits_zero() -> None:
    res = runner.invoke(app, ["template", "inspect"])
    assert res.exit_code == 0, res.output
    assert "canto=entrada" in res.output
    assert "oe=III" in res.output
    assert "liturgy.mystery_of_faith" in res.output
    assert "0 erro(s)" in res.output


def test_inspect_reports_errors_and_exits_one(tmp_path: Path) -> None:
    presentation = Presentation(str(default_template_path()))
    presentation.slides[0].notes_slide.notes_text_frame.text = "@pallotti canto=xyz"
    broken = tmp_path / "broken.pptx"
    presentation.save(str(broken))

    res = runner.invoke(app, ["template", "inspect", "--template", str(broken)])
    assert res.exit_code == 1
    assert "ERROR [slide 1]" in res.output
    assert "momento de canto desconhecido" in res.output


def test_inspect_missing_file_exits_one(tmp_path: Path) -> None:
    res = runner.invoke(app, ["template", "inspect", "--template", str(tmp_path / "x.pptx")])
    assert res.exit_code == 1
    assert "não encontrado" in res.output


def test_help_of_template_commands() -> None:
    for args in (
        ["template", "--help"],
        ["template", "inspect", "--help"],
        ["template", "preview", "--help"],
    ):
        res = runner.invoke(app, args)
        assert res.exit_code == 0
    assert "--output" in runner.invoke(app, ["template", "preview", "--help"]).output


def test_tools_report_matches_template() -> None:
    report = TemplateTools().inspect()
    assert len(report.rows) == len(MassTemplate.load().slides)
    assert report.prayers == ["III"]
    assert not report.has_errors


@pytest.mark.soffice
def test_preview_renders_one_png_per_slide(tmp_path: Path) -> None:
    res = runner.invoke(app, ["template", "preview", "--output", str(tmp_path / "previa")])
    assert res.exit_code == 0, res.output
    pngs = sorted((tmp_path / "previa").glob("slide-*.png"))
    assert len(pngs) == len(MassTemplate.load().slides)
    assert all(p.stat().st_size > 1000 for p in pngs)
