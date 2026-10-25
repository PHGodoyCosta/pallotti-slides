"""Builds src/pallotti_slides/resources/template/mass_template.pptx from the example deck.

Reads ``examples/Missa - 26° Dom do T.C - 26.09.2026 - OE III.pptx`` (read-only) and, following
docs/MODELO_PPTX.md:

* keeps ONE prototype slide per song moment and replaces its texts by ``{{canto.rotulo}}`` and
  ``{{canto.texto}}``;
* writes the ``@pallotti ...`` tag on the first line of each slide's notes;
* replaces the variable texts by ``{{...}}`` tags (inside the original run, so the formatting
  stays);
* adds the communion / post-communion / recessional songs, by copying existing slides.

Nothing is drawn by code: slides are only kept, removed, copied and reordered.
Every slide number below is a slide of the EXAMPLE (1-based). The values marked PROVISIONAL are
decisions waiting for Pedro's approval; change them here and run the script again.

Usage: uv run python scripts/create_template_from_example.py
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide

from pallotti_slides.domain.enums import SongMoment
from pallotti_slides.domain.models import Mass
from pallotti_slides.presentation.slides_xml import duplicate_slide, keep_slides
from pallotti_slides.presentation.template import (
    MOMENT_TAGS,
    SECOND_READING_GROUP,
    SONG_GROUP_PREFIX,
    TAG_PREFIX,
    MassTemplate,
)

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = ROOT / "examples" / "Missa - 26° Dom do T.C - 26.09.2026 - OE III.pptx"
FIXTURE = ROOT / "tests" / "fixtures" / "masses" / "26tc_2026.json"
OUTPUT = ROOT / "src" / "pallotti_slides" / "resources" / "template" / "mass_template.pptx"

# --- Song moments of the example -------------------------------------------------------------
# The prototype is the first slide of each song block; the other slides of the block go away.
PROTOTYPES: dict[SongMoment, int] = {
    SongMoment.MANTRA: 3,
    SongMoment.ENTRANCE: 5,
    SongMoment.PENITENTIAL_ACT: 10,
    SongMoment.GLORIA: 15,
    SongMoment.ACCLAMATION: 26,
    SongMoment.OFFERTORY: 36,
}
# Separator image that precedes the song and disappears together with it.
SEPARATORS: dict[SongMoment, int] = {
    SongMoment.ENTRANCE: 4,
    SongMoment.PENITENTIAL_ACT: 9,
    SongMoment.GLORIA: 14,
    SongMoment.OFFERTORY: 35,
}
SONG_SLIDES_TO_DROP = {6, 7, 8, 11, 12, 13, 16, 17, 18, 19, 37, 38, 39}

# Slides that go away together when there is no second reading.
SECOND_READING_SLIDES = (24, 25)

# --- Eucharistic prayer block (PROVISIONAL) --------------------------------------------------
PRAYER_ID = "III"
PRAYER_SLIDES = range(47, 61)  # 47..60

# --- End of the Mass (PROVISIONAL) -----------------------------------------------------------
# Before the final slide 69 come, in this order, a separator + prototype for each moment below.
FINAL_SLIDE = 69
NEW_SONGS = (SongMoment.COMMUNION, SongMoment.POST_COMMUNION, SongMoment.RECESSIONAL)
NEW_SEPARATOR_SOURCE = 4  # Jesus image: always the separator image
NEW_PROTOTYPE_SOURCE = 36  # offertory prototype

# --- Text tags: (example slide, field of the reference Mass, tag in the template) ------------
# The old text must exist, exactly, in ONE run of the slide (sanity check against the fixture).
TEXT_TAGS: tuple[tuple[int, str, str], ...] = (
    (2, "celebration.title", "celebracao.titulo"),
    (21, "liturgy.first_reading", "liturgia.primeira_leitura"),
    (23, "liturgy.psalm_number", "liturgia.salmo_numero"),
    (23, "liturgy.psalm_refrain", "liturgia.salmo_refrao"),
    (24, "liturgy.second_reading", "liturgia.segunda_leitura"),
    (28, "liturgy.gospel", "liturgia.evangelho"),
    (31, "liturgy.intercessions_response", "liturgia.preces_refrao"),
    (50, "liturgy.mystery_of_faith", "liturgia.misterio_fe"),
)


def fail(message: str) -> None:
    sys.exit(f"ERRO: {message}")


def text_shapes(slide: Slide) -> list[Shape]:
    return [s for s in slide.shapes if isinstance(s, Shape) and s.has_text_frame]


def replace_run_text(slide: Slide, number: int, old: str, new: str) -> None:
    """Replace the text of the single run that holds exactly ``old`` (formatting is kept)."""
    runs = [
        run
        for shape in text_shapes(slide)
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.text == old
    ]
    if len(runs) != 1:
        fail(f"slide {number}: esperava 1 trecho com {old!r}, achei {len(runs)}")
    runs[0].text = new


def make_song_prototype(slide: Slide, number: int, moment: SongMoment) -> None:
    """Label -> ``{{canto.rotulo}}``; body -> one paragraph with ``{{canto.texto}}``."""
    shapes = text_shapes(slide)
    labels = [s for s in shapes if s.text_frame.text.strip() == moment.label]
    if len(labels) != 1:
        fail(f"slide {number}: rótulo {moment.label!r} não encontrado (achei {len(labels)})")
    label = labels[0]
    bodies = [s for s in shapes if s is not label and s.text_frame.text.strip()]
    if len(bodies) != 1:
        fail(f"slide {number}: esperava 1 corpo de canto, achei {len(bodies)}")
    body = bodies[0]

    label.text_frame.paragraphs[0].runs[0].text = "{{canto.rotulo}}"

    paragraphs = body.text_frame.paragraphs
    first = paragraphs[0]
    for paragraph in paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    for run in first.runs[1:]:
        run._r.getparent().remove(run._r)
    first.runs[0].text = "{{canto.texto}}"


def set_tag(slide: Slide, tag: str) -> None:
    slide.notes_slide.notes_text_frame.text = f"{TAG_PREFIX} {tag}"


def main() -> None:
    if not EXAMPLE.is_file():
        fail(f"exemplo não encontrado: {EXAMPLE}")
    mass = Mass.model_validate_json(FIXTURE.read_text(encoding="utf-8"))
    values: dict[str, str] = {"celebration.title": mass.celebration.title}
    for name, value in mass.liturgy.model_dump().items():
        if isinstance(value, str):
            values[f"liturgy.{name}"] = value

    presentation = Presentation(str(EXAMPLE))
    original = list(presentation.slides)
    if len(original) != FINAL_SLIDE:
        fail(f"o exemplo deveria ter {FINAL_SLIDE} slides, tem {len(original)}")

    def slide_of(number: int) -> Slide:
        return original[number - 1]

    # 1. Texts: song prototypes and variable fields (before copying, so copies inherit them).
    for moment, number in PROTOTYPES.items():
        make_song_prototype(slide_of(number), number, moment)
    for number, field, tag in TEXT_TAGS:
        replace_run_text(slide_of(number), number, values[field], f"{{{{{tag}}}}}")

    # 2. Tags for the slides that stay from the example.
    tags: dict[int, str] = {number: "fixo" for number in range(1, FINAL_SLIDE + 1)}
    for moment, number in PROTOTYPES.items():
        tags[number] = f"canto={MOMENT_TAGS[moment]}"
    for moment, number in SEPARATORS.items():
        tags[number] = f"grupo={SONG_GROUP_PREFIX}{MOMENT_TAGS[moment]}"
    for number in SECOND_READING_SLIDES:
        tags[number] = f"grupo={SECOND_READING_GROUP}"
    for number in PRAYER_SLIDES:
        tags[number] = f"oe={PRAYER_ID}"

    # 3. Final order: example slides (minus the dropped ones) + the new songs before slide 69.
    ordered: list[tuple[Slide, str]] = []
    for number in range(1, FINAL_SLIDE + 1):
        if number in SONG_SLIDES_TO_DROP:
            continue
        if number == FINAL_SLIDE:
            for moment in NEW_SONGS:
                group = f"{SONG_GROUP_PREFIX}{MOMENT_TAGS[moment]}"
                separator = duplicate_slide(presentation, slide_of(NEW_SEPARATOR_SOURCE))
                prototype = duplicate_slide(presentation, slide_of(NEW_PROTOTYPE_SOURCE))
                ordered.append((separator, f"grupo={group}"))
                ordered.append((prototype, f"canto={MOMENT_TAGS[moment]}"))
        ordered.append((slide_of(number), tags[number]))

    # 4. Keep/reorder/rename, then write the tags in the notes.
    keep_slides(presentation, [slide for slide, _ in ordered])
    for slide, tag in ordered:
        set_tag(slide, tag)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(OUTPUT))
    print(f"Modelo gravado em {OUTPUT} ({len(ordered)} slides)")

    # 5. Check the result with the same code the program uses.
    problems = MassTemplate.load(OUTPUT).validate()
    for notice in problems:
        print(f"{notice.level.upper()}: {notice.field or '-'}: {notice.message}")
    if any(notice.level == "error" for notice in problems):
        sys.exit(1)


if __name__ == "__main__":
    main()
