"""Generates tests/fixtures/masses/26tc_2026.json from the example presentation.

Reads ``examples/Missa - 26° Dom do T.C - 26.09.2026 - OE III.pptx`` with python-pptx and
copies every text from it (nothing is written from memory). The slide numbers below follow
docs/MODELO_PPTX.md §3. The result is reviewed by hand and versioned.

Usage: uv run python scripts/generate_example_fixture.py
"""

import datetime as dt
import json
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide

from pallotti_slides.domain.enums import (
    EucharisticPrayer,
    LiturgicalSeason,
    SongMoment,
    SongSource,
    StanzaType,
)
from pallotti_slides.domain.models import Celebration, Line, Liturgy, Mass, Song, Stanza
from pallotti_slides.domain.text import normalize_title

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = ROOT / "examples" / "Missa - 26° Dom do T.C - 26.09.2026 - OE III.pptx"
OUTPUT = ROOT / "tests" / "fixtures" / "masses" / "26tc_2026.json"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
TITLE_SLIDE = 2
FIRST_READING_SLIDE = 21
PSALM_SLIDE = 23
SECOND_READING_SLIDE = 24
GOSPEL_SLIDE = 28
INTERCESSIONS_SLIDE = 31
MYSTERY_OF_FAITH_SLIDE = 50
# Song slides (1-based), per moment. Moment label is read from each slide and checked.
SONG_SLIDES: dict[SongMoment, range] = {
    SongMoment.MANTRA: range(3, 4),
    SongMoment.ENTRANCE: range(5, 9),
    SongMoment.PENITENTIAL_ACT: range(10, 14),
    SongMoment.GLORIA: range(15, 20),
    SongMoment.ACCLAMATION: range(26, 27),
    SongMoment.OFFERTORY: range(36, 40),
}
BODY_SHAPE = "PlaceHolder 1"


def _clean(text: str) -> str:
    return text.replace("\xa0", " ").strip()


def _shape(slide: Slide, name: str) -> Shape:
    for shape in slide.shapes:
        if shape.name == name and shape.has_text_frame:
            return shape  # type: ignore[return-value]
    raise LookupError(f"forma {name!r} não encontrada no slide")


def _label_of(slide: Slide) -> str:
    """Moment label: the small text box that is not the body."""
    labels = [
        _clean(shape.text_frame.text)
        for shape in slide.shapes
        if shape.has_text_frame and shape.name != BODY_SHAPE and shape.text_frame.text.strip()
    ]
    if len(labels) != 1:
        raise LookupError(f"rótulo ambíguo: {labels}")
    return labels[0]


def _body_text(slide: Slide) -> str:
    return _shape(slide, BODY_SHAPE).text_frame.text


def _read_lines(slide: Slide) -> list[Line]:
    """Lines of the body: one per paragraph / line break; bold from the runs."""
    lines: list[Line] = []
    for paragraph in _shape(slide, BODY_SHAPE).text_frame.paragraphs:
        pieces: list[list[tuple[str, bool]]] = [[]]
        runs = {run._r: run for run in paragraph.runs}
        for child in paragraph._p:
            name = etree.QName(child).localname
            if name == "r":
                pieces[-1].append((runs[child].text, runs[child].font.bold is True))
            elif name == "br":
                pieces.append([])
        for piece in pieces:
            visible = [bold for text, bold in piece if text.strip()]
            text = _clean("".join(text for text, _ in piece))
            lines.append(Line(text=text, bold=bool(visible) and all(visible)))
    while lines and not lines[0].text:
        lines.pop(0)
    while lines and not lines[-1].text:
        lines.pop()
    return lines


def _build_song(presentation_slides: list[Slide], moment: SongMoment, numbers: range) -> Song:
    stanzas: list[Stanza] = []
    order: list[int] = []
    seen: dict[tuple[tuple[str, bool], ...], int] = {}
    for number in numbers:
        slide = presentation_slides[number - 1]
        label = _label_of(slide)
        if label != moment.label:
            raise ValueError(f"slide {number}: rótulo {label!r} != {moment.label!r}")
        lines = _read_lines(slide)
        key = tuple((line.text, line.bold) for line in lines)
        if key not in seen:
            stanza_id = len(stanzas) + 1
            kind = StanzaType.CHORUS if all(line.bold for line in lines) else StanzaType.VERSE
            stanzas.append(Stanza(id=stanza_id, type=kind, lines=lines))
            seen[key] = stanza_id
        order.append(seen[key])
    return Song(
        title=stanzas[0].lines[0].text,  # provisional: the pptx has no song title
        artist=None,
        source=SongSource.MANUAL,
        source_url=None,
        stanzas=stanzas,
        order=order,
    )


def _all_text(slide: Slide) -> str:
    return "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


def build_mass() -> Mass:
    presentation = Presentation(str(EXAMPLE))
    slides = list(presentation.slides)

    def slide(number: int) -> Slide:
        return slides[number - 1]

    title = _clean(_all_text(slide(TITLE_SLIDE)))
    if normalize_title(title) != title:
        raise ValueError(f"título fora do padrão: {title!r}")

    psalm_lines = [
        _clean(p.text) for p in _shape(slide(PSALM_SLIDE), BODY_SHAPE).text_frame.paragraphs
    ]
    psalm_lines = [line for line in psalm_lines if line]
    psalm_number = psalm_lines[0].removeprefix("Salmo ").strip()
    preces = _clean(_all_text(slide(INTERCESSIONS_SLIDE)).split("\n")[0]).removeprefix("R:").strip()

    file_match = re.search(r"- (\d{2})\.(\d{2})\.(\d{4}) - OE (I{1,3}|IV)\.pptx$", EXAMPLE.name)
    if file_match is None:
        raise ValueError("nome do exemplo fora do padrão")
    day, month, year = (int(g) for g in file_match.groups()[:3])

    songs = {
        moment: _build_song(slides, moment, numbers) for moment, numbers in SONG_SLIDES.items()
    }
    return Mass(
        celebration=Celebration(
            title=title,
            short_title=None,
            date=dt.date(year, month, day),
            season=LiturgicalSeason.ORDINARY_TIME,
            year=None,
        ),
        liturgy=Liturgy(
            first_reading=_clean(_all_text(slide(FIRST_READING_SLIDE))),
            psalm_number=psalm_number,
            psalm_refrain=psalm_lines[1],
            second_reading=_clean(_all_text(slide(SECOND_READING_SLIDE))),
            gospel=_clean(_all_text(slide(GOSPEL_SLIDE))),
            intercessions_response=preces,
            eucharistic_prayer=EucharisticPrayer(file_match.group(4)),
            mystery_of_faith=_clean(_all_text(slide(MYSTERY_OF_FAITH_SLIDE))),
        ),
        has_gloria=True,
        songs=songs,
    )


def summary(mass: Mass) -> str:
    rows = []
    for moment, song in mass.songs.items():
        kinds = ",".join(f"{s.id}:{s.type.value}" for s in song.stanzas)
        rows.append(
            f"  {moment.value:16} estrofes={len(song.stanzas)} ordem={song.order} [{kinds}]"
            f" título provisório={song.title!r}"
        )
    return "\n".join(rows)


def main() -> int:
    mass = build_mass()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    data = mass.model_dump(mode="json")
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Fixture gravada em {OUTPUT}")
    print(summary(mass))
    return 0


if __name__ == "__main__":
    sys.exit(main())
