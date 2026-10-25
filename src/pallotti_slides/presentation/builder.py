"""Executes a ``SlidePlan``: template + prayer block + song decks -> deck (ARQUITETURA §5)."""

from pathlib import Path

from pptx import Presentation
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide

from pallotti_slides.presentation.errors import TemplateError
from pallotti_slides.presentation.planner import SlidePlan, TextContent
from pallotti_slides.presentation.slides_xml import (
    BodyLine,
    SlideImporter,
    SlideImportError,
    duplicate_slide,
    keep_slides,
    remove_note_lines,
    replace_body,
    replace_run_tags,
    text_tags,
)
from pallotti_slides.presentation.template import FIELD_BY_TAG, TAG_BY_FIELD, TAG_PREFIX


def _open(path: Path) -> PresentationType:
    try:
        return Presentation(str(path))
    except Exception as error:  # python-pptx raises many kinds of errors on bad files
        raise TemplateError(f"Não foi possível abrir {path.name}: {error}") from error


class PptxBuilder:
    """Writes the deck: assembles the slides and fills the ``{{...}}`` tags.

    1. a template slide is the slide itself (first use) or a copy made from the untouched
       template slide (next uses); a slide of another file (prayer block, cached song deck) is
       imported (``slides_xml.SlideImporter``); unused template slides are removed;
    2. tags are replaced at the run level (formatting kept); a song body still in the plan is
       rebuilt as one centred paragraph per line (``slides_xml.replace_body``);
    3. the ``@pallotti`` lines are removed from the speaker notes.
    No file read here is ever modified.
    """

    def build(self, template_path: Path, plan: SlidePlan, destination: Path) -> Path:
        presentation = _open(template_path)
        originals = list(presentation.slides)
        others: dict[Path, list[Slide]] = {}
        for target in plan:
            if target.deck is not None and target.deck not in others:
                if not target.deck.is_file():
                    raise TemplateError(f"Arquivo de slides não encontrado: {target.deck}")
                others[target.deck] = list(_open(target.deck).slides)
            available = originals if target.deck is None else others[target.deck]
            if not 0 <= target.source < len(available):
                where = "no modelo" if target.deck is None else f"em {target.deck.name}"
                raise TemplateError(
                    f"O plano cita o slide {target.source + 1}, que não existe {where}."
                )

        # 1. Allocate the slides first, so every copy comes from an untouched template slide.
        importer = SlideImporter(presentation)
        used: set[int] = set()
        slides: list[Slide] = []
        for target in plan:
            if target.deck is not None:
                try:
                    slides.append(importer.import_slide(others[target.deck][target.source]))
                except SlideImportError as error:
                    raise TemplateError(
                        f"Não foi possível usar os slides de {target.deck.name}: {error}"
                    ) from error
                continue
            original = originals[target.source]
            if target.source in used:
                slides.append(duplicate_slide(presentation, original))
            else:
                used.add(target.source)
                slides.append(original)

        # 2. Texts.
        for slide, target in zip(slides, plan, strict=True):
            where = f"de {target.deck.name}" if target.deck is not None else "do modelo"
            self._fill(slide, target.replacements, f"{target.source + 1} {where}")

        # 3. Order, removal and notes.
        keep_slides(presentation, slides)
        for slide in slides:
            if slide.has_notes_slide:
                remove_note_lines(slide.notes_slide.element, TAG_PREFIX)

        # Write next to the destination and rename: a failure never leaves half a file.
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_name(f".{destination.stem}.partial.pptx")
        try:
            presentation.save(str(temporary))
            temporary.replace(destination)
        finally:
            temporary.unlink(missing_ok=True)
        return destination

    @staticmethod
    def _fill(slide: Slide, replacements: dict[str, TextContent], where: str) -> None:
        root = slide.element
        for field, content in replacements.items():
            if content.is_body:
                tag = TAG_BY_FIELD[field]
                lines = [BodyLine(line.text, line.bold) for line in content.lines]
                try:
                    replace_body(root, tag, lines, content.size_pt or 0)
                except ValueError as error:
                    raise TemplateError(f"O slide {where} não tem {{{{{tag}}}}}.") from error

        def value(tag: str) -> str:
            field = FIELD_BY_TAG.get(tag)
            content = replacements.get(field) if field else None
            if content is None or content.is_body:
                raise TemplateError(f"A etiqueta {{{{{tag}}}}} do slide {where} não tem valor.")
            return content.text

        replace_run_tags(root, value)
        left = text_tags(root)
        if left:
            raise TemplateError(f"Etiquetas sem substituição no slide {where}: {left}.")
