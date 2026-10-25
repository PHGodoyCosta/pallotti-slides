"""Song deck cache (M3b): one .pptx per song, made from the prototype and the pagination.

``<directory>/<key>.pptx`` holds one slide per page of the song (the moment's prototype with
the body already filled and ``{{canto.rotulo}}`` still in place); ``<key>.json`` keeps the page
count and whether a line overflows. The key is a hash of everything that changes those
slides: the stanzas in execution order, the ``TypographyConfig``, the measuring font and the
prototype slide itself (its XML and images). So a song is regenerated automatically when its
lyrics, the typography or the template prototype change, and the same file serves every
moment whose prototype is identical (only the label changes, when the deck is assembled).
"""

import dataclasses
import hashlib
import json
import os
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide

from pallotti_slides.domain.enums import SongMoment
from pallotti_slides.domain.models import Song
from pallotti_slides.presentation.errors import TemplateError
from pallotti_slides.presentation.paginator import SongPage, SongPaginator
from pallotti_slides.presentation.slides_xml import (
    BodyLine,
    duplicate_slide,
    keep_slides,
    remove_note_lines,
    replace_body,
)
from pallotti_slides.presentation.template import (
    SONG_TEXT_FIELD,
    TAG_BY_FIELD,
    TAG_PREFIX,
    MassTemplate,
)

# Bump when the layout of the cached decks changes (forces every song to be regenerated).
CACHE_VERSION = 1


@dataclass(frozen=True)
class SongDeck:
    """A cached song: ``path`` has ``pages`` slides, in order (``pages`` > 0)."""

    path: Path
    pages: int
    overflow: bool = False  # some line does not fit the body box even in ``min_font``


class SongDeckCache:
    """Gives the cached deck of a song, creating it when needed (``directory`` is injectable)."""

    def __init__(self, directory: Path, template: MassTemplate, paginator: SongPaginator) -> None:
        self.directory = directory
        self._template = template
        self._paginator = paginator
        self._prototype_keys: dict[SongMoment, str] = {}
        self.generated: list[Path] = []  # decks written by this instance (diagnostics, tests)

    def key(self, moment: SongMoment, song: Song) -> str:
        """Cache key of ``song`` in the prototype of ``moment``."""
        stanzas = [
            {"type": stanza.type.value, "lines": [[line.text, line.bold] for line in stanza.lines]}
            for stanza in song.sequence()
        ]
        data = {
            "version": CACHE_VERSION,
            "stanzas": stanzas,
            "typography": dataclasses.asdict(self._paginator.config),
            "substitute_font": self._paginator.measurer.uses_substitute,
            "prototype": self._prototype_key(moment),
        }
        encoded = json.dumps(data, ensure_ascii=False, sort_keys=True).encode("utf-8")
        return hashlib.sha256(encoded).hexdigest()[:32]

    def deck(self, moment: SongMoment, song: Song) -> SongDeck | None:
        """The deck of ``song`` for ``moment``; ``None`` when the song has no lines."""
        if not any(stanza.lines for stanza in song.sequence()):
            return None
        key = self.key(moment, song)
        path = self.directory / f"{key}.pptx"
        cached = self._read(path)
        if cached is not None:
            return cached
        area = self._template.song_area(moment)
        body = self._template.body_format(moment)
        pages = self._paginator.paginate(song, area, body)
        if not pages:
            return None
        measurer = self._paginator.measurer
        overflow = not all(measurer.fits(p.lines, p.size_pt, area, body) for p in pages)
        self._write(moment, pages, path)
        meta = {"pages": len(pages), "overflow": overflow}
        _atomic_write(path.with_suffix(".json"), json.dumps(meta).encode("utf-8"))
        self.generated.append(path)
        return SongDeck(path=path, pages=len(pages), overflow=overflow)

    def _read(self, path: Path) -> SongDeck | None:
        meta = path.with_suffix(".json")
        if not (path.is_file() and meta.is_file()):
            return None
        try:
            data = json.loads(meta.read_text(encoding="utf-8"))
            pages, overflow = int(data["pages"]), bool(data["overflow"])
        except (OSError, ValueError, KeyError, TypeError):
            return None
        return SongDeck(path=path, pages=pages, overflow=overflow) if pages > 0 else None

    def _prototype_key(self, moment: SongMoment) -> str:
        """Hash of the prototype slide of ``moment``: XML (without notes), layout and images."""
        known = self._prototype_keys.get(moment)
        if known is not None:
            return known
        prototype = self._template.prototype(moment)
        if prototype is None:
            raise TemplateError(f"O modelo não tem o protótipo de canto {moment.label}.")
        slide = self._open_template().slides[prototype.index]
        digest = hashlib.sha256(etree.tostring(slide.element, method="c14n"))
        digest.update(slide.slide_layout.name.encode("utf-8"))
        for rel_id in sorted(slide.part.rels):
            rel = slide.part.rels[rel_id]
            if rel.is_external:
                digest.update(f"{rel_id}:{rel.target_ref}".encode())
            elif rel.reltype not in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
                digest.update(rel_id.encode() + hashlib.sha1(rel.target_part.blob).digest())
        digest.update(f"{self._template.slide_width}x{self._template.slide_height}".encode())
        self._prototype_keys[moment] = digest.hexdigest()
        return self._prototype_keys[moment]

    def _open_template(self) -> PresentationType:
        try:
            return Presentation(str(self._template.path))
        except Exception as error:  # python-pptx raises many kinds of errors on bad files
            raise TemplateError(
                f"Não foi possível abrir o modelo {self._template.path.name}: {error}"
            ) from error

    def _write(self, moment: SongMoment, pages: list[SongPage], path: Path) -> None:
        prototype = self._template.prototype(moment)
        assert prototype is not None  # checked by song_area()
        presentation = self._open_template()
        original: Slide = presentation.slides[prototype.index]
        slides = [original] + [duplicate_slide(presentation, original) for _ in pages[1:]]
        tag = TAG_BY_FIELD[SONG_TEXT_FIELD]
        for slide, page in zip(slides, pages, strict=True):
            lines = [BodyLine(line.text, line.bold) for line in page.lines]
            replace_body(slide.element, tag, lines, page.size_pt)
        keep_slides(presentation, slides)
        if original.has_notes_slide:
            remove_note_lines(original.notes_slide.element, TAG_PREFIX)
        self.directory.mkdir(parents=True, exist_ok=True)
        temporary = path.with_name(f".{path.stem}.{os.getpid()}.partial.pptx")
        try:
            presentation.save(str(temporary))
            temporary.replace(path)
        finally:
            temporary.unlink(missing_ok=True)


def _atomic_write(path: Path, data: bytes) -> None:
    temporary = path.with_name(f".{path.name}.{os.getpid()}.partial")
    try:
        temporary.write_bytes(data)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
