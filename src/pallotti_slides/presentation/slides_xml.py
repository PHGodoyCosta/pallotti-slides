"""Duplicate, remove and reorder slides and replace tagged text at the XML level.

python-pptx cannot duplicate slides and ``text_frame.text = ...`` loses the formatting, so the
generator works on the XML with lxml (namespaces preserved).
"""

import re
from collections.abc import Callable, Sequence
from copy import deepcopy
from dataclasses import dataclass

from lxml.etree import _Element
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.presentation import Presentation
from pptx.slide import Slide

_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_SKIPPED_RELATIONS = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}
_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_CSLD = f"{{{_P_NS}}}cSld"
_SP_TREE = f"{{{_P_NS}}}spTree"


def _copy_common_slide_data(source: _Element, target: _Element) -> None:
    """Copy ``p:cSld`` (background + shape tree), keeping the target's ``p:spTree`` element."""
    target_tree = target.find(_SP_TREE)
    for child in list(target):
        if child is not target_tree:
            target.remove(child)
    for child in list(target_tree):
        target_tree.remove(child)
    after = target_tree
    seen_tree = False
    for child in source:
        if child.tag == _SP_TREE:
            seen_tree = True
            for name, value in child.attrib.items():
                target_tree.set(name, value)
            for grandchild in child:
                target_tree.append(deepcopy(grandchild))
        elif not seen_tree:
            target_tree.addprevious(deepcopy(child))
        else:
            after.addnext(deepcopy(child))
            after = after.getnext()


def duplicate_slide(presentation: Presentation, slide: Slide) -> Slide:
    """Append a copy of ``slide`` (same layout, same images) at the end of the deck.

    The copy gets its own slide part; images and other targets are shared with the original
    through new relationships. Notes are not copied.
    """
    copy: Slide = presentation.slides.add_slide(slide.slide_layout)
    new_part = copy.part

    renamed: dict[str, str] = {}
    for rel_id, rel in slide.part.rels.items():
        if rel.reltype in _SKIPPED_RELATIONS:
            continue
        if rel.is_external:
            renamed[rel_id] = new_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            renamed[rel_id] = new_part.relate_to(rel.target_part, rel.reltype)

    # Edit in place (not replace the root): python-pptx caches ``copy.shapes`` on the old tree.
    source_root = slide.element
    target_root = copy.element
    for name, value in source_root.attrib.items():
        target_root.set(name, value)
    for child in list(target_root):
        if child.tag != _CSLD:
            target_root.remove(child)
    for child in source_root:
        if child.tag != _CSLD:
            target_root.append(deepcopy(child))
    _copy_common_slide_data(source_root.find(_CSLD), target_root.find(_CSLD))

    prefix = f"{{{_R_NS}}}"
    for element in target_root.iter():
        for name, value in list(element.attrib.items()):
            if name.startswith(prefix) and value in renamed:
                element.set(name, renamed[value])
    return copy


def keep_slides(presentation: Presentation, ordered: Sequence[Slide]) -> None:
    """Keep only ``ordered``, in that order; every other slide is removed from the deck.

    Removed slide parts become unreachable and are not written on save. Slide parts are
    renamed ``slide1.xml``… following the final order.
    """
    id_list = presentation.slides._sldIdLst
    by_id = {entry.id: entry for entry in id_list}
    wanted = [slide.slide_id for slide in ordered]
    for entry in list(id_list):
        if entry.id not in wanted:
            presentation.part.drop_rel(entry.rId)
        id_list.remove(entry)
    for slide_id in wanted:
        id_list.append(by_id[slide_id])
    presentation.part.rename_slide_parts([entry.rId for entry in id_list])


# --- Text at the run level (M3) -----------------------------------------------------------------

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_A_P = f"{{{_A_NS}}}p"
_A_R = f"{{{_A_NS}}}r"
_A_T = f"{{{_A_NS}}}t"
_A_PPR = f"{{{_A_NS}}}pPr"
_A_RPR = f"{{{_A_NS}}}rPr"
_A_END = f"{{{_A_NS}}}endParaRPr"
_TAG = re.compile(r"\{\{(.*?)\}\}")


@dataclass(frozen=True)
class BodyLine:
    """One line (paragraph) of a song body."""

    text: str
    bold: bool


def text_tags(element: _Element) -> list[str]:
    """Names of every ``{{tag}}`` still written in the runs under ``element``."""
    return [name.strip() for t in element.iter(_A_T) for name in _TAG.findall(t.text or "")]


def replace_run_tags(element: _Element, values: Callable[[str], str]) -> None:
    """Replace each ``{{tag}}`` inside its run by ``values(tag)``; run formatting is kept."""
    for t in element.iter(_A_T):
        text = t.text or ""
        if "{{" in text:
            t.text = _TAG.sub(lambda match: values(match.group(1).strip()), text)


def find_tag_paragraph(element: _Element, tag: str) -> _Element | None:
    """The ``a:p`` whose runs hold ``{{tag}}``."""
    marker = f"{{{{{tag}}}}}"
    for run in element.iter(_A_R):
        t = run.find(_A_T)
        if t is not None and marker in (t.text or ""):
            return run.getparent()
    return None


def _set_size_and_bold(run_props: _Element, size_pt: int, bold: bool) -> None:
    run_props.set("sz", str(size_pt * 100))
    run_props.set("b", "1" if bold else "0")


def replace_body(element: _Element, tag: str, lines: Sequence[BodyLine], size_pt: int) -> None:
    """Replace the paragraph holding ``{{tag}}`` by one centred paragraph per line.

    Each new paragraph copies the ``a:pPr`` of the prototype paragraph (with ``algn="ctr"``)
    and its run copies the ``a:rPr`` of the tag run, with ``sz`` = ``size_pt`` and ``b`` from
    the line. Nothing else of the shape changes.
    """
    paragraph = find_tag_paragraph(element, tag)
    if paragraph is None:
        raise ValueError(f"{{{{{tag}}}}} not found")
    marker = f"{{{{{tag}}}}}"
    proto_run = next(run for run in paragraph.iter(_A_R) if marker in (run.findtext(_A_T) or ""))
    proto_ppr = paragraph.find(_A_PPR)
    proto_end = paragraph.find(_A_END)
    anchor = paragraph
    for line in lines or [BodyLine("", False)]:
        new = deepcopy(paragraph)
        for child in list(new):
            new.remove(child)
        ppr = deepcopy(proto_ppr) if proto_ppr is not None else new.makeelement(_A_PPR, {})
        ppr.set("algn", "ctr")
        new.append(ppr)
        if line.text:
            run = deepcopy(proto_run)
            run_props = run.find(_A_RPR)
            if run_props is None:
                run_props = run.makeelement(_A_RPR, {})
                run.insert(0, run_props)
            _set_size_and_bold(run_props, size_pt, line.bold)
            t = run.find(_A_T)
            assert t is not None
            t.text = line.text
            new.append(run)
        if proto_end is not None:
            end = deepcopy(proto_end)
            _set_size_and_bold(end, size_pt, line.bold)
            new.append(end)
        anchor.addnext(new)
        anchor = new
    parent = paragraph.getparent()
    assert parent is not None
    parent.remove(paragraph)


def remove_note_lines(notes: _Element, prefix: str) -> None:
    """Remove from a notes slide every paragraph whose text starts with ``prefix``.

    A text body left without paragraphs gets one empty paragraph (required by the schema).
    """
    for paragraph in list(notes.iter(_A_P)):
        text = "".join(t.text or "" for t in paragraph.iter(_A_T))
        if text.strip().startswith(prefix):
            body = paragraph.getparent()
            assert body is not None
            body.remove(paragraph)
            if body.find(_A_P) is None:
                body.append(body.makeelement(_A_P, {}))
