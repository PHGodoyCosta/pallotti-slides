"""The tagged PowerPoint template (MODELO_PPTX.md): read slide tags and text tags.

Tags in the template are in Portuguese (the team edits them in PowerPoint); this module maps
them to the English enums and field names of the code.
"""

import re
from dataclasses import dataclass, field
from enum import StrEnum
from pathlib import Path
from typing import Literal, Self

from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.slide import Slide

from pallotti_slides.domain.enums import EucharisticPrayer, SongMoment
from pallotti_slides.domain.models import Notice
from pallotti_slides.presentation.errors import TemplateError
from pallotti_slides.presentation.measurer import Box

TAG_PREFIX = "@pallotti"
SECOND_READING_GROUP = "segunda_leitura"
SONG_GROUP_PREFIX = "canto:"

# Tag moment id (Portuguese, as written in the template) <-> SongMoment.
MOMENT_BY_TAG: dict[str, SongMoment] = {
    "mantra": SongMoment.MANTRA,
    "entrada": SongMoment.ENTRANCE,
    "ato_penitencial": SongMoment.PENITENTIAL_ACT,
    "gloria": SongMoment.GLORIA,
    "aclamacao": SongMoment.ACCLAMATION,
    "ofertorio": SongMoment.OFFERTORY,
    "comunhao": SongMoment.COMMUNION,
    "pos_comunhao": SongMoment.POST_COMMUNION,
    "saida": SongMoment.RECESSIONAL,
}
MOMENT_TAGS: dict[SongMoment, str] = {moment: tag for tag, moment in MOMENT_BY_TAG.items()}

# ``{{tag}}`` written in the slides (Portuguese) -> field name used by the code (English).
FIELD_BY_TAG: dict[str, str] = {
    "celebracao.titulo": "celebration.title",
    "liturgia.primeira_leitura": "liturgy.first_reading",
    "liturgia.salmo_numero": "liturgy.psalm_number",
    "liturgia.salmo_refrao": "liturgy.psalm_refrain",
    "liturgia.segunda_leitura": "liturgy.second_reading",
    "liturgia.evangelho": "liturgy.gospel",
    "liturgia.preces_refrao": "liturgy.intercessions_response",
    "liturgia.misterio_fe": "liturgy.mystery_of_faith",
    "canto.rotulo": "song.label",
    "canto.texto": "song.text",
}
TAG_BY_FIELD: dict[str, str] = {name: tag for tag, name in FIELD_BY_TAG.items()}

SONG_LABEL_FIELD = "song.label"
SONG_TEXT_FIELD = "song.text"
MYSTERY_FIELD = "liturgy.mystery_of_faith"

Level = Literal["info", "warning", "error"]
_TEXT_TAG = re.compile(r"\{\{(.*?)\}\}")
_PRAYER_BY_ID = {prayer.value: prayer for prayer in EucharisticPrayer}


class TagKind(StrEnum):
    """Role of a slide in the template."""

    FIXED = "fixed"
    GROUP = "group"
    SONG = "song"
    PRAYER = "prayer"
    INVALID = "invalid"


@dataclass(frozen=True)
class SlideTag:
    """The ``@pallotti ...`` tag from the first line of a slide's notes.

    ``group`` is the group id as written in the template (``canto:entrada``,
    ``segunda_leitura``); song prototypes belong to ``canto:<moment>``.
    """

    kind: TagKind
    raw: str = ""
    explicit: bool = True
    group: str | None = None
    moment: SongMoment | None = None
    prayer: EucharisticPrayer | None = None
    error: str | None = None

    @classmethod
    def parse(cls, first_line: str | None) -> Self:
        """Parse the first line of the notes. Never raises: problems go to ``error``."""
        line = (first_line or "").strip()
        if not line.startswith(TAG_PREFIX):
            return cls(kind=TagKind.FIXED, raw=line, explicit=False)
        body = line[len(TAG_PREFIX) :].strip()
        if body == "fixo":
            return cls(kind=TagKind.FIXED, raw=line)
        key, sep, value = body.partition("=")
        key, value = key.strip(), value.strip()
        if sep and value and " " not in body:
            if key == "grupo":
                return cls._group(line, value)
            if key == "canto":
                moment = MOMENT_BY_TAG.get(value)
                if moment is not None:
                    return cls(
                        kind=TagKind.SONG,
                        raw=line,
                        group=f"{SONG_GROUP_PREFIX}{value}",
                        moment=moment,
                    )
                return cls._invalid(line, f"momento de canto desconhecido: '{value}'")
            if key == "oe":
                prayer = _PRAYER_BY_ID.get(value)
                if prayer is not None:
                    return cls(kind=TagKind.PRAYER, raw=line, prayer=prayer)
                return cls._invalid(line, f"Oração Eucarística desconhecida: '{value}'")
        return cls._invalid(line, "etiqueta mal formada")

    @classmethod
    def _group(cls, line: str, group: str) -> Self:
        known = group == SECOND_READING_GROUP or (
            group.startswith(SONG_GROUP_PREFIX)
            and group.removeprefix(SONG_GROUP_PREFIX) in MOMENT_BY_TAG
        )
        if not known:
            return cls._invalid(line, f"grupo desconhecido: '{group}'")
        moment = MOMENT_BY_TAG.get(group.removeprefix(SONG_GROUP_PREFIX))
        return cls(kind=TagKind.GROUP, raw=line, group=group, moment=moment)

    @classmethod
    def _invalid(cls, line: str, error: str) -> Self:
        return cls(kind=TagKind.INVALID, raw=line, error=error)


@dataclass(frozen=True)
class TemplateSlide:
    """One slide of the template. ``index`` is 0-based (``MassTemplate.slides[index]``)."""

    index: int
    tag: SlideTag
    fields: frozenset[str]
    text_box: Box | None = None
    unknown_tags: tuple[str, ...] = field(default=())
    split_tags: tuple[str, ...] = field(default=())


def default_template_path() -> Path:
    """Location of the template shipped with the program."""
    return Path(__file__).resolve().parents[1] / "resources" / "template" / "mass_template.pptx"


class MassTemplate:
    """The tagged template, read once into memory (the file is never modified)."""

    def __init__(self, path: Path, slides: list[TemplateSlide]) -> None:
        self.path = path
        self.slides = slides

    @classmethod
    def load(cls, path: Path | None = None) -> Self:
        path = path or default_template_path()
        if not path.is_file():
            raise TemplateError(f"Modelo do PowerPoint não encontrado: {path}")
        try:
            presentation = Presentation(str(path))
        except Exception as error:  # python-pptx raises many kinds of errors on bad files
            raise TemplateError(f"Não foi possível abrir o modelo {path.name}: {error}") from error
        slides = [_read_slide(index, slide) for index, slide in enumerate(presentation.slides)]
        return cls(path, slides)

    def available_prayers(self) -> list[EucharisticPrayer]:
        """Eucharistic prayers that have a block in the template, in numeric order."""
        found = {s.tag.prayer for s in self.slides if s.tag.prayer is not None}
        return [prayer for prayer in EucharisticPrayer if prayer in found]

    def validate(self) -> list[Notice]:
        """Problems of the template (MODELO_PPTX.md §6)."""
        notices: list[Notice] = []

        def add(level: Level, slide: TemplateSlide | None, message: str) -> None:
            where = f"slide {slide.index + 1}" if slide else None
            notices.append(Notice(level=level, message=message, field=where))

        prototypes: dict[SongMoment, list[TemplateSlide]] = {}
        for slide in self.slides:
            if slide.tag.kind is TagKind.INVALID:
                add("error", slide, f"Etiqueta inválida ({slide.tag.error}): {slide.tag.raw!r}.")
            elif not slide.tag.explicit:
                add(
                    "warning",
                    slide,
                    "Slide sem etiqueta @pallotti nas anotações (tratado como fixo).",
                )
            if slide.tag.kind is TagKind.SONG and slide.tag.moment is not None:
                prototypes.setdefault(slide.tag.moment, []).append(slide)
            for tag in slide.unknown_tags:
                add("error", slide, f"Etiqueta de texto desconhecida: {{{{{tag}}}}}.")
            for tag in slide.split_tags:
                add(
                    "error",
                    slide,
                    f"Etiqueta {{{{{tag}}}}} dividida em vários trechos de texto; "
                    "apague e digite de novo de uma vez.",
                )
            is_song = slide.tag.kind is TagKind.SONG
            if is_song and SONG_TEXT_FIELD not in slide.fields:
                add("error", slide, "Protótipo de canto sem {{canto.texto}}.")
            if is_song and SONG_LABEL_FIELD not in slide.fields:
                add("warning", slide, "Protótipo de canto sem {{canto.rotulo}}.")
            song_fields = slide.fields & {SONG_LABEL_FIELD, SONG_TEXT_FIELD}
            if song_fields and not is_song:
                add("error", slide, "Etiquetas {{canto.*}} só podem ficar em protótipo de canto.")

        for moment in SongMoment:
            found = prototypes.get(moment, [])
            tag = MOMENT_TAGS[moment]
            if not found:
                add("error", None, f"Falta o protótipo de canto '{tag}' ({moment.label}).")
            for extra in found[1:]:
                add("error", extra, f"Protótipo de canto '{tag}' duplicado.")

        for prayer in self.available_prayers():
            block = [s for s in self.slides if s.tag.prayer is prayer]
            if not any(MYSTERY_FIELD in s.fields for s in block):
                add(
                    "warning",
                    None,
                    f"O bloco da OE {prayer.value} não tem {{{{liturgia.misterio_fe}}}}.",
                )
        if not self.available_prayers():
            add(
                "error",
                None,
                "O modelo não tem nenhum bloco de Oração Eucarística (@pallotti oe=...).",
            )
        return notices


def _first_notes_line(slide: Slide) -> str | None:
    if not slide.has_notes_slide:
        return None
    text = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
    for line in text.splitlines():
        if line.strip():
            return line
    return None


def _shape_box(shape: BaseShape) -> Box | None:
    if None in (shape.left, shape.top, shape.width, shape.height):
        return None
    return Box(
        left=int(shape.left or 0),
        top=int(shape.top or 0),
        width=int(shape.width or 0),
        height=int(shape.height or 0),
    )


def _read_slide(index: int, slide: Slide) -> TemplateSlide:
    tag = SlideTag.parse(_first_notes_line(slide))
    fields: set[str] = set()
    unknown: list[str] = []
    split: list[str] = []
    text_box: Box | None = None
    for shape in slide.shapes:
        if not isinstance(shape, Shape) or not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            run_texts = [run.text for run in paragraph.runs]
            whole = _TEXT_TAG.findall("".join(run_texts))
            whole_in_run = [name for text in run_texts for name in _TEXT_TAG.findall(text)]
            for name in whole_in_run:
                known = FIELD_BY_TAG.get(name.strip())
                if known is None:
                    unknown.append(name)
                    continue
                fields.add(known)
                if known == SONG_TEXT_FIELD and text_box is None:
                    text_box = _shape_box(shape)
            if len(whole) != len(whole_in_run):
                split.extend(name for name in whole if name not in whole_in_run)
    return TemplateSlide(
        index=index,
        tag=tag,
        fields=frozenset(fields),
        text_box=text_box,
        unknown_tags=tuple(unknown),
        split_tags=tuple(split),
    )
